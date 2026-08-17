"""Generate the project final presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FILE = PROJECT_ROOT / "Final_Presentation.pptx"
NAVY = RGBColor(18, 42, 74)
BLUE = RGBColor(33, 113, 181)
DARK = RGBColor(40, 40, 40)


def add_title(slide, title, subtitle=""):
    title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.4), Inches(12), Inches(0.7))
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.68), Inches(1.12), Inches(11.5), Inches(0.45))
        subtitle_paragraph = subtitle_box.text_frame.paragraphs[0]
        subtitle_paragraph.text = subtitle
        subtitle_paragraph.font.size = Pt(14)
        subtitle_paragraph.font.color.rgb = DARK


def add_bullets(slide, items, top=1.75):
    box = slide.shapes.add_textbox(Inches(0.95), Inches(top), Inches(11.3), Inches(4.8))
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(14)
    return box


def add_footer(slide, number):
    line = slide.shapes.add_shape(1, Inches(0.65), Inches(7.05), Inches(12.0), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()
    box = slide.shapes.add_textbox(Inches(11.8), Inches(7.12), Inches(0.5), Inches(0.25))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = str(number)
    paragraph.alignment = PP_ALIGN.RIGHT
    paragraph.font.size = Pt(10)
    paragraph.font.color.rgb = DARK


def add_slide(presentation, title, bullets, subtitle=""):
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, title, subtitle)
    add_bullets(slide, bullets)
    add_footer(slide, len(presentation.slides))


def main():
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(0.9), Inches(2.15), Inches(11.5), Inches(1.0))
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = "Jira Ticket Quality and Duplicate Analysis"
    paragraph.font.size = Pt(34)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    subtitle = slide.shapes.add_textbox(Inches(0.95), Inches(3.3), Inches(10.5), Inches(0.6))
    subtitle_paragraph = subtitle.text_frame.paragraphs[0]
    subtitle_paragraph.text = "Final project presentation"
    subtitle_paragraph.font.size = Pt(20)
    subtitle_paragraph.font.color.rgb = BLUE
    add_footer(slide, 1)

    add_slide(presentation, "Project objective", [
        "Assess the completeness and rule compliance of exported Jira tickets.",
        "Identify likely duplicate tickets using only historical tickets.",
        "Produce clear terminal feedback and an Excel report for reviewers.",
    ])
    add_slide(presentation, "Solution architecture", [
        "Input: Jira export workbook.",
        "Validation: modular checks in checks/validators.py.",
        "Duplicate analysis: date-aware candidate selection and text similarity.",
        "Reporting: quality KPIs, selected ticket, and similarity report in Excel.",
    ])
    add_slide(presentation, "Validation logic", [
        "WARNING: information is missing or needs review; the ticket remains valid.",
        "ERROR: a required rule is not respected; the ticket is Not Valid.",
        "Class_002 requires Label_017 or Label_008; violations are errors.",
        "Closed tickets require a valid resolution; missing resolution is an error.",
    ])
    add_slide(presentation, "Duplicate-analysis safeguards", [
        "Only tickets created before the selected ticket are compared.",
        "A missing or invalid creation date stops duplicate analysis explicitly.",
        "The report records NOT ANALYSED with the reason instead of implying no duplicates.",
        "Similarity uses sentence embeddings when available, with TF-IDF fallback.",
    ])
    add_slide(presentation, "Quality checks refined", [
        "Attachment evidence requires semantic terms such as attachment, log, or trace.",
        "Jira markup characters ! and [^ are not treated as evidence.",
        "TC validation recognizes only exact TC_<number> identifiers.",
        "Generic words such as test and SW_BUG labels do not create TC false positives.",
    ])
    add_slide(presentation, "Deliverables", [
        "Complete Python source code with modular validation, analysis, and reporting components.",
        "README with prerequisites, installation, run commands, outputs, and tests.",
        "Excel report with Quality KPI Dataset, Quality KPIs, Selected Ticket, and Similarity Report sheets.",
        "Automated unit tests covering validation rules and duplicate-analysis safeguards.",
    ])
    add_slide(presentation, "How to run", [
        "Install dependencies: python -m pip install -r requirements.txt",
        "Run analysis: python main\\main.py",
        "Run tests: python -m unittest discover -s tests -v",
        "Generate this deck: python tools\\create_presentation.py",
    ])
    add_slide(presentation, "Conclusion", [
        "The solution provides consistent ticket-quality decisions and transparent review signals.",
        "Mandatory rule violations invalidate tickets; missing information remains visible as warnings.",
        "Duplicate results are chronologically safe and clearly state when analysis cannot be performed.",
    ])

    presentation.save(OUTPUT_FILE)
    print(f"Presentation created: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
